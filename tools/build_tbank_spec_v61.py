"""Build TBANK validator spec v6.1 — v6.0 body + production supplement, no legacy-disable."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

SRC = Path(__file__).with_name("_v6_spec_from_docx.txt")
OUT_DOCX = Path(r"c:\Users\fanis\Downloads\TBANK_VALIDATOR_SPEC_v6_1.docx")
OUT_MD = Path(r"c:\Users\fanis\Downloads\TBANK_VALIDATOR_SPEC_v6_1.md")

# Drop entire sections (header line starts section; ends before next N. or Приложение)
DROP_SECTIONS = frozenset({
    "1. Обязательная замена старой логики",
    "18. Критерии приёмки полной замены",
})

DROP_LINE_SUBSTR = (
    "Не объединять этот документ со старой логикой",
    "Отключить и удалить прежние правила",
    "ПРИМЕНИТЬ TBANK VALIDATOR MASTER SPEC",
    "ПОЛНОСТЬЮ УДАЛИТЬ",
    "Старый decision engine Т-Банка отключён",
    "ПОЛНАЯ ЗАМЕНА ПРОЦЕССА ВАЛИДАЦИИ",
    "MASTER SPECIFICATION v6.0 — единый самодостаточный документ",
    "ДИРЕКТИВА С НАИВЫСШИМ ПРИОРИТЕТОМ",
    "1. Обязательная замена старой логики",
    "18. Критерии приёмки",
    "Финальная директива для копирования в ИИ/валидатор",
    "Не сливать со старыми правилами",
    "Полностью удалить старый Т-Банк decision engine",
)

SUPPLEMENT = r"""
================================================================================
РАЗДЕЛ 19–30. PRODUCTION-РЕАЛИЗАЦИЯ (дополнение к v6.0, 12.07.2026)
================================================================================

19. Архитектура runtime

19.1. Точки входа
  • detector/__init__.py → route(pdf_bytes)
  • detector/tbank.py → detector/tbank_v6/engine.py::analyze()
  • VALIDATOR_VERSION = 6.0.0, engine = tbank_v6

19.2. Global Preflight (до tbank_v6)
  Модуль: detector/hardening_v2/global_preflight.py
  Выполняется для Т-Банка до bank analyzer. Может вернуть ФЕЙК:
  • G-FONT-005 GLOBAL_TEXT_GLYPH_RENDER_MISMATCH
  • G-GRAPH-003 G_GRAPH_003_LATENT_CONFLICT
  • G-STATUS-002 STATUS_INTERNAL_CONFLICT
  • G-SEM-001 AMOUNT_ARITHMETIC_MISMATCH
  • Structural: JAVASCRIPT, EMBEDDED, XFA, xref/EOF corruption
  Stats: details.global_preflight

19.3. Фактический порядок стадий tbank_v6 (detector/tbank_v6/stages.py)
  1. intake          — SHA-256, лимит 8 MB
  2. raw_preflight   — structure, xref, incremental
  3. streams         — compression/Length
  4. active_content  — JS/XFA/embedded
  5. content_ast     — BT/ET, edit trace, known skeletons
  6. layout_fingerprint — T-TBANK-TEXT-LAYOUT-FINGERPRINT-001
  7. fonts_cmap_glyph   — K-FONT-001/002, pdf_forensics
  8. metadata_generation — K-TBANK-KEYWORDS-GENERATION-001
  9. semantic_geometry  — подметод, SBP, anti_edit
  10. differential_parity — fitz SBP-ID parity
  11. cross_document_intelligence — anti_edit DB

20. K-TBANK-KEYWORDS-GENERATION-001

  Модуль: detector/tbank_keywords_generation.py
  Hard: TBANK_KEYWORDS_GENERATION_MISMATCH (tier A)
  Scope: Subject /reports/IB/Receipt

  • Third token /Keywords — поколение backend (pipe-separated tail).
  • До 10.07.2026 MSK: tail «991» — OK (legacy).
  • С 10.07.2026: ожидается «DOCS-2035».
  • formed_at ≥ TRANSITION + tail «991» → hard ФЕЙК.
  • Неизвестный новый tail без mismatch → diagnostic (не auto-ban).

21. T-TBANK-TEXT-LAYOUT-FINGERPRINT-001

  Модуль: detector/tbank_text_layout_fingerprint.py
  Hard: TBANK_TEXT_LAYOUT_FINGERPRINT (tier A)
  Scope: все подметоды Т-Банка

  • Raw content stream → end_x value-строк (F1, x≥100).
  • R ≈ median(right edge) ≈ 250 pt.
  • Hard: overflow > 0.05 pt И fitz подтверждает (dual-engine).
  • Mode: active.

22. K-TBANK-SBP-CONTENT-002 + K-TBANK-SBP-GEOMETRY-001

  22.1 Content — detector/tbank_sbp_content.py::validate_tbank_sbp_id()
  Только channel=SBP.

  Hard codes:
  • SBP_CIPHER_MISSING, SBP_CIPHER_STRUCTURE, SBP_CIPHER_TIMESTAMP
  • Profile block ID[22:26]=«0011» (только Jasper/OpenPDF)

  Timestamp decode (строже generic sbp_cipher):
  • year_digit=ID[1], doy=ID[2:5], UTC HMS=ID[5:11]
  • MSK date from receipt → UTC −3h
  • ±1 min, ±2 sec на границе минуты

  Diagnostic: class ID[17:19] + suffix ID[26:32] (01+760501)

  22.2 Geometry — detector/tbank_sbp_geometry.py
  Hard: TBANK_SBP_GEOMETRY_MISMATCH

  • Raw SBP lines + fitz match by text
  • R = median F1 value end_x (235–265, default 250)
  • Hard: overflow/shortfall > 0.05 pt + fitz confirms
  • Skip if opid parity failed

  22.3 Generic sbp_cipher.py — не используется в v6 pipeline

23. Anti-edit (detector/anti_edit.py, SQLite anti_edit.db)

  OPERATION_ID_REUSED        tier A — тот же ID, другие amount/date/receiver
  RECEIPT_TEXT_LAYER_MISSING tier A — пустой text + content
  PDF_MODDATE_EDITED         tier B — ModDate drift > 60 sec

24. Дополнительные semantic hard (stages.py)

  TEXT_LAYER_INCONSISTENT, BROKEN_CYRILLIC_MAPPING, FIELD_FORMAT_INVALID,
  TEXT_EXTRACTION_MAPPING_ANOMALY, TBANK_CONTENT_STREAM_EDIT, F3_NOT_ALSRUBL

25. Known-fake signatures

  K-FONT-001 — font rebuilder A combo (tbank_font_rebuilder.py)
  K-FONT-002 — F2 glyf/loca exact SHA (0/128 originals)
  KNOWN_GENERATOR_SKELETONS — 6 skeleton hashes

26. Verdict и UI

  Verdict: «ЧИСТО» / «ФЕЙК» (expert report: «ОРИГИНАЛ»)
  Score: 0 / 60 (interface compatibility)
  Expert: detector/tbank_v6/explain.py → privileged users in bot

27. Enroll (не verdict)

  detector/tbank_enroll.py — кнопка «оригинал»: skeleton, FF2, render rows в корпус

28. Declared HARD_CODES не wired в stages

  AMOUNT_MISMATCH_BETWEEN_TOTAL_AND_DETAILS, FIELD_ORDER_MISMATCH,
  FIELD_SET_MISMATCH, MISSING_REQUIRED_FIELD_BLOCK, SBP_CIPHER_REFERENCE

29. Regression anchors (v6)

  REG-O-001: 128 originals corpus → ЧИСТО
  REG-O-002: Receipt (6).pdf DOCS-2035 → ЧИСТО
  REG-O-003: receipt_10.07.2026 (28).pdf 991 post-transition → ФЕЙК
  REG-O-004: receipt_12.07.2026 (4).pdf → ЧИСТО

30. Карта модулей

  tbank_v6/engine.py, stages.py, verdict.py, rules.py, known_signatures.py, explain.py
  tbank_keywords_generation.py, tbank_text_layout_fingerprint.py
  tbank_sbp_content.py, tbank_sbp_geometry.py, tbank_font_rebuilder.py
  anti_edit.py, hardening_v2/global_preflight.py, global_text_glyph_render.py
  pdf_forensics.py, structure.py

================================================================================
"""


def _section_header(line: str) -> str | None:
    s = line.strip()
    m = re.match(r"^(\d+\.\s.+)$", s)
    if m:
        return m.group(1)
    if s.startswith("Приложение "):
        return s
    return None


def filter_body() -> list[str]:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    out: list[str] = []
    skip = False
    skip_directive = False
    current_section: str | None = None

    for line in lines:
        if line.strip().startswith("Финальная директива"):
            skip_directive = True
            continue
        if skip_directive:
            continue

        hdr = _section_header(line)
        if hdr:
            current_section = hdr
            if any(hdr.startswith(d.split(".", 1)[0] + ".") and d in hdr for d in DROP_SECTIONS):
                skip = True
                continue
            if hdr in DROP_SECTIONS or any(d in hdr for d in DROP_SECTIONS):
                skip = True
                continue
            skip = False

        if skip:
            continue
        if any(x in line for x in DROP_LINE_SUBSTR):
            continue
        if line.strip().startswith("- ") and "стары" in line.lower():
            continue
        if "ПОЛОСЫ РЕНДЕРА УДАЛИТЬ" in line:
            out.append("Render rows не участвуют в verdict (corpus enroll — отдельный процесс).")
            continue
        if "Значение third token (991, DOCS-2035 или новое) является opaque" in line:
            out.append(
                "Значение third token профилирует поколение документа; "
                "см. §20 K-TBANK-KEYWORDS-GENERATION-001."
            )
            continue
        out.append(line)

    return out


def build_document() -> str:
    header = """TBANK VALIDATOR — MASTER SPECIFICATION v6.1
PDF-квитанции Т-Банка — все подметоды (СБП, телефон, карта)
Production · engine: detector/tbank_v6 · VALIDATOR_VERSION 6.0.0
Дата: 12.07.2026

v6.0 (11.07.2026) + дополнения production-реализации.
Без разделов про отключение legacy — описан только активный pipeline.

"""
    body = "\n".join(filter_body())
    # insert supplement before first Приложение
    parts = body.split("Приложение A.")
    if len(parts) == 2:
        body = parts[0].rstrip() + "\n\n" + SUPPLEMENT + "\n\nПриложение A." + parts[1]
    else:
        body = body + "\n\n" + SUPPLEMENT
    return header + body


def write_docx(text: str) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    t = doc.add_heading("TBANK VALIDATOR — MASTER SPECIFICATION v6.1", 0)
    t.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph("Production · tbank_v6 · 12.07.2026").alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph("")

    for line in text.splitlines():
        s = line.strip()
        if not s:
            doc.add_paragraph("")
            continue
        if s.startswith("===="):
            continue
        if re.match(r"^\d+\.\s", s) and len(s) < 120 and s[0].isdigit():
            doc.add_heading(s, level=2)
        elif s.startswith("Приложение "):
            doc.add_heading(s, level=1)
        elif re.match(r"^\d+\.\d+\.", s):
            doc.add_heading(s, level=3)
        else:
            doc.add_paragraph(line)

    doc.save(OUT_DOCX)


def main() -> None:
    text = build_document()
    OUT_MD.write_text(text, encoding="utf-8")
    write_docx(text)
    print(OUT_MD)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
